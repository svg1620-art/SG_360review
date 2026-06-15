"""Движок отчёта по оцениваемому (разделы 7–8 ТЗ). Видит только админ.

Агрегат строится строго по правилам анонимности:
- self / manager — не анонимны, показываются всегда;
- peer / subordinate — отдельные столбцы только при достижении порога;
- «другие» = пул анонимных источников (peer + subordinate); показывается, только
  если суммарно ответивших >= порога И это не позволяет вычесть скрытую подгруппу.
"""
import random
from collections import defaultdict
from io import BytesIO

from flask import (
    Blueprint,
    abort,
    flash,
    g,
    redirect,
    render_template,
    send_file,
    url_for,
)
from openpyxl import Workbook
from openpyxl.chart import LineChart, RadarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor

from .auth import login_required
from .db import get_db
from .limits import is_pro

bp = Blueprint("report", __name__)

BLIND_DELTA = 1.0  # порог Δ для слепых/скрытых зон (шкала 1–5)
TOP_N = 3          # сколько компетенций в «сильные»/«рост»


def _others_visible(p, s, threshold):
    """Можно ли показать композит «другие» (peer+sub) без утечки скрытой подгруппы."""
    if p + s < threshold:
        return False
    # Утечка: одна анонимная группа показана (>=порога), другая непуста, но скрыта (<порога).
    if (p >= threshold and 0 < s < threshold) or (s >= threshold and 0 < p < threshold):
        return False
    return True


def _mean(pair):
    return pair[0] / pair[1] if pair[1] else None


def _rating_aggregate(conn, cycle_id, subject_id, threshold):
    """Средние по компетенциям и источникам + флаги видимости. Переиспользуется для динамики."""
    with conn.cursor() as cur:
        cur.execute(
            "SELECT competency_name FROM cycle_questions WHERE cycle_id = %s ORDER BY id",
            (cycle_id,),
        )
        comp_order, seen = [], set()
        for r in cur.fetchall():
            if r["competency_name"] not in seen:
                seen.add(r["competency_name"])
                comp_order.append(r["competency_name"])

        cur.execute(
            """
            SELECT relation, count(*) AS n FROM assignments
            WHERE cycle_id = %s AND subject_id = %s AND status = 'submitted'
            GROUP BY relation
            """,
            (cycle_id, subject_id),
        )
        counts = {row["relation"]: row["n"] for row in cur.fetchall()}

        cur.execute(
            """
            SELECT a.relation, cq.competency_name, r.rating
            FROM responses r
            JOIN assignments a ON a.id = r.assignment_id
            JOIN cycle_questions cq ON cq.id = r.cycle_question_id
            WHERE a.cycle_id = %s AND a.subject_id = %s AND a.status = 'submitted'
              AND cq.qtype = 'rating' AND r.rating IS NOT NULL
            """,
            (cycle_id, subject_id),
        )
        rating_rows = cur.fetchall()

    p, s = counts.get("peer", 0), counts.get("subordinate", 0)
    peer_visible, sub_visible = p >= threshold, s >= threshold
    others_visible = _others_visible(p, s, threshold)

    acc = defaultdict(lambda: defaultdict(lambda: [0, 0]))  # comp -> relation -> [sum, n]
    others = defaultdict(lambda: [0, 0])                    # comp -> [sum, n] для peer+sub
    for row in rating_rows:
        comp, rel, val = row["competency_name"], row["relation"], row["rating"]
        acc[comp][rel][0] += val
        acc[comp][rel][1] += 1
        if rel in ("peer", "subordinate"):
            others[comp][0] += val
            others[comp][1] += 1

    means = {}
    for comp in comp_order:
        src = acc.get(comp, {})
        means[comp] = {
            "self": _mean(src["self"]) if "self" in src else None,
            "manager": _mean(src["manager"]) if "manager" in src else None,
            "peer": _mean(src["peer"]) if ("peer" in src and peer_visible) else None,
            "subordinate": _mean(src["subordinate"]) if ("subordinate" in src and sub_visible) else None,
            "others": _mean(others[comp]) if (comp in others and others_visible) else None,
        }
    return {
        "means": means,
        "counts": counts,
        "competencies": comp_order,
        "peer_visible": peer_visible,
        "sub_visible": sub_visible,
        "others_visible": others_visible,
    }


def _open_answers(conn, cycle_id, subject_id, peer_visible, sub_visible):
    """Открытые ответы, сгруппированные по вопросу; peer/sub — анонимно и в перемешку."""
    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, competency_name, question_text FROM cycle_questions WHERE cycle_id = %s AND qtype = 'open' ORDER BY id",
            (cycle_id,),
        )
        questions = cur.fetchall()
        cur.execute(
            """
            SELECT a.relation, r.cycle_question_id AS qid, r.text_answer
            FROM responses r
            JOIN assignments a ON a.id = r.assignment_id
            JOIN cycle_questions cq ON cq.id = r.cycle_question_id
            WHERE a.cycle_id = %s AND a.subject_id = %s AND a.status = 'submitted'
              AND cq.qtype = 'open' AND r.text_answer IS NOT NULL
            """,
            (cycle_id, subject_id),
        )
        rows = cur.fetchall()

    by_q = defaultdict(lambda: {"self": [], "manager": [], "peers": [], "subs": []})
    for row in rows:
        bucket = by_q[row["qid"]]
        rel, text = row["relation"], row["text_answer"]
        if rel == "self":
            bucket["self"].append(text)
        elif rel == "manager":
            bucket["manager"].append(text)
        elif rel == "peer" and peer_visible:
            bucket["peers"].append(text)
        elif rel == "subordinate" and sub_visible:
            bucket["subs"].append(text)

    blocks = []
    for q in questions:
        b = by_q.get(q["id"], {"self": [], "manager": [], "peers": [], "subs": []})
        random.shuffle(b["peers"])  # без привязки к автору, в перемешанном порядке
        random.shuffle(b["subs"])
        if b["self"] or b["manager"] or b["peers"] or b["subs"]:
            blocks.append(
                {
                    "competency": q["competency_name"],
                    "question": q["question_text"],
                    "self": b["self"],
                    "manager": b["manager"],
                    "peers": b["peers"],
                    "subs": b["subs"],
                }
            )
    return blocks


def _dynamics(conn, company_id, employee_id, threshold, comp_order):
    """Средний балл «другие» по компетенциям от закрытого цикла к закрытому циклу."""
    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT c.id, c.title, c.period_label, cs.id AS subject_id
            FROM cycles c
            JOIN cycle_subjects cs ON cs.cycle_id = c.id
            WHERE c.company_id = %s AND cs.employee_id = %s AND c.status = 'closed'
            ORDER BY c.starts_at NULLS LAST, c.id
            """,
            (company_id, employee_id),
        )
        cycles = cur.fetchall()
    if len(cycles) < 2:
        return None

    labels = [c["period_label"] or c["title"] for c in cycles]
    series = {comp: [] for comp in comp_order}
    for c in cycles:
        agg = _rating_aggregate(conn, c["id"], c["subject_id"], threshold)
        for comp in comp_order:
            val = agg["means"].get(comp, {}).get("others") if agg else None
            series[comp].append(round(val, 2) if val is not None else None)
    # Оставляем только компетенции, где есть хотя бы одна точка.
    series = {k: v for k, v in series.items() if any(x is not None for x in v)}
    return {"labels": labels, "series": series}


def _load_target(db, cycle_id, subject_id):
    """Загружает цикл и оцениваемого с проверкой принадлежности компании."""
    with db.cursor() as cur:
        cur.execute(
            "SELECT * FROM cycles WHERE id = %s AND company_id = %s",
            (cycle_id, g.company_id),
        )
        cycle = cur.fetchone()
        if cycle is None:
            abort(404)
        cur.execute(
            """
            SELECT cs.employee_id, e.full_name
            FROM cycle_subjects cs JOIN employees e ON e.id = cs.employee_id
            WHERE cs.id = %s AND cs.cycle_id = %s
            """,
            (subject_id, cycle_id),
        )
        subj = cur.fetchone()
        if subj is None:
            abort(404)
        cur.execute("SELECT anon_threshold FROM companies WHERE id = %s", (g.company_id,))
        threshold = cur.fetchone()["anon_threshold"]
    return cycle, subj, threshold


def _assemble(db, cycle, subject_id, employee_id, threshold, with_dynamics):
    """Собирает все данные отчёта. Общий источник для HTML и XLSX."""
    cycle_id = cycle["id"]
    agg = _rating_aggregate(db, cycle_id, subject_id, threshold)
    means, comp_order = agg["means"], agg["competencies"]

    radar_comps = [c for c in comp_order if means[c]["self"] is not None and means[c]["others"] is not None]
    radar = {
        "labels": radar_comps,
        "self": [round(means[c]["self"], 2) for c in radar_comps],
        "others": [round(means[c]["others"], 2) for c in radar_comps],
        "manager": [round(means[c]["manager"], 2) if means[c]["manager"] is not None else None for c in radar_comps],
    }
    has_manager_line = any(v is not None for v in radar["manager"])

    others_pairs = [(c, means[c]["others"]) for c in comp_order if means[c]["others"] is not None]
    strong = sorted(others_pairs, key=lambda x: -x[1])[:TOP_N]
    growth = sorted(others_pairs, key=lambda x: x[1])[:TOP_N]
    blind, hidden = [], []
    for c in comp_order:
        sv, ov = means[c]["self"], means[c]["others"]
        if sv is None or ov is None:
            continue
        if sv - ov > BLIND_DELTA:
            blind.append((c, sv, ov, sv - ov))
        elif ov - sv > BLIND_DELTA:
            hidden.append((c, sv, ov, ov - sv))
    blind.sort(key=lambda x: -x[3])
    hidden.sort(key=lambda x: -x[3])

    open_blocks = _open_answers(db, cycle_id, subject_id, agg["peer_visible"], agg["sub_visible"])
    dynamics = _dynamics(db, g.company_id, employee_id, threshold, comp_order) if with_dynamics else None

    return {
        "agg": agg,
        "means": means,
        "competencies": comp_order,
        "radar": radar,
        "has_manager_line": has_manager_line,
        "zones": {"strong": strong, "growth": growth, "blind": blind, "hidden": hidden},
        "open_blocks": open_blocks,
        "dynamics": dynamics,
    }


@bp.route("/cycles/<int:cycle_id>/subjects/<int:subject_id>/report")
@login_required
def report(cycle_id, subject_id):
    db = get_db()
    cycle, subj, threshold = _load_target(db, cycle_id, subject_id)
    if cycle["status"] != "closed":
        flash("Отчёт доступен после закрытия цикла.")
        return redirect(url_for("cycles.detail", cycle_id=cycle_id))

    # Сравнение динамики — только на Pro (раздел 9 ТЗ).
    ctx = _assemble(db, cycle, subject_id, subj["employee_id"], threshold, with_dynamics=is_pro())
    agg = ctx["agg"]
    return render_template(
        "report/report.html",
        cycle=cycle,
        subject_id=subject_id,
        subject_name=subj["full_name"],
        threshold=threshold,
        counts=agg["counts"],
        competencies=ctx["competencies"],
        means=ctx["means"],
        peer_visible=agg["peer_visible"],
        sub_visible=agg["sub_visible"],
        others_visible=agg["others_visible"],
        radar=ctx["radar"],
        has_manager_line=ctx["has_manager_line"],
        zones=ctx["zones"],
        open_blocks=ctx["open_blocks"],
        dynamics=ctx["dynamics"],
    )


def _two_cell_anchor(from_col, from_row, to_col, to_row):
    """TwoCellAnchor (а не oneCellAnchor — иначе ломается парсер SG, см. CLAUDE.md)."""
    anchor = TwoCellAnchor()
    anchor._from = AnchorMarker(col=from_col, row=from_row)
    anchor.to = AnchorMarker(col=to_col, row=to_row)
    return anchor


def _xlsx_cell(value, visible=True):
    if not visible:
        return "скрыто"
    return round(value, 2) if value is not None else ""


def _build_workbook(cycle, subject_name, threshold, ctx):
    means, comps, agg = ctx["means"], ctx["competencies"], ctx["agg"]
    wb = Workbook()

    # Лист 1: сводка по компетенциям
    ws = wb.active
    ws.title = "Сводка"
    ws.append([f"Отчёт 360°: {subject_name}"])
    ws.append([f"Цикл: {cycle['title']}", f"Порог анонимности: {threshold}"])
    ws.append([])
    ws.append(["Компетенция", "Самооценка", "Руководитель", "Коллеги", "Подчинённые", "Другие"])
    for c in comps:
        m = means[c]
        ws.append([
            c,
            _xlsx_cell(m["self"]),
            _xlsx_cell(m["manager"]),
            _xlsx_cell(m["peer"], agg["peer_visible"]),
            _xlsx_cell(m["subordinate"], agg["sub_visible"]),
            _xlsx_cell(m["others"], agg["others_visible"]),
        ])

    # Лист 2: я vs другие + радар
    ws2 = wb.create_sheet("Я vs другие")
    header = ["Компетенция", "Самооценка", "Другие"]
    if ctx["has_manager_line"]:
        header.append("Руководитель")
    ws2.append(header)
    labels = ctx["radar"]["labels"]
    for c in labels:
        row = [c, round(means[c]["self"], 2), round(means[c]["others"], 2)]
        if ctx["has_manager_line"]:
            mv = means[c]["manager"]
            row.append(round(mv, 2) if mv is not None else None)
        ws2.append(row)
    if labels:
        n = len(labels)
        last_col = 4 if ctx["has_manager_line"] else 3
        chart = RadarChart()
        chart.title = "Я vs другие"
        data = Reference(ws2, min_col=2, min_row=1, max_col=last_col, max_row=n + 1)
        cats = Reference(ws2, min_col=1, min_row=2, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 1, 5
        chart.anchor = _two_cell_anchor(last_col + 1, 1, last_col + 9, 20)
        ws2.add_chart(chart)

    # Лист 3: зоны
    ws3 = wb.create_sheet("Зоны")
    ws3.append(["Сильные стороны (по «другие»)"])
    for c, v in ctx["zones"]["strong"]:
        ws3.append([c, round(v, 2)])
    ws3.append([])
    ws3.append(["Зоны роста (по «другие»)"])
    for c, v in ctx["zones"]["growth"]:
        ws3.append([c, round(v, 2)])
    ws3.append([])
    ws3.append(["Слепые зоны (я ≫ другие)", "самооценка", "другие", "Δ"])
    for c, sv, ov, d in ctx["zones"]["blind"]:
        ws3.append([c, round(sv, 2), round(ov, 2), round(d, 2)])
    ws3.append([])
    ws3.append(["Скрытые сильные (другие ≫ я)", "самооценка", "другие", "Δ"])
    for c, sv, ov, d in ctx["zones"]["hidden"]:
        ws3.append([c, round(sv, 2), round(ov, 2), round(d, 2)])

    # Лист 4: открытые ответы (с учётом анонимности)
    ws4 = wb.create_sheet("Открытые ответы")
    for b in ctx["open_blocks"]:
        ws4.append([b["competency"]])
        ws4.append([b["question"]])
        for t in b["self"]:
            ws4.append(["самооценка", t])
        for t in b["manager"]:
            ws4.append(["руководитель", t])
        for t in b["peers"]:
            ws4.append(["коллеги (аноним.)", t])
        for t in b["subs"]:
            ws4.append(["подчинённые (аноним.)", t])
        ws4.append([])

    # Лист 5: динамика
    dyn = ctx["dynamics"]
    if dyn and dyn["series"]:
        ws5 = wb.create_sheet("Динамика")
        comp_names = list(dyn["series"].keys())
        ws5.append(["Цикл"] + comp_names)
        for i, label in enumerate(dyn["labels"]):
            ws5.append([label] + [dyn["series"][cn][i] for cn in comp_names])
        n, m = len(dyn["labels"]), len(comp_names)
        chart = LineChart()
        chart.title = "Динамика «другие»"
        data = Reference(ws5, min_col=2, min_row=1, max_col=m + 1, max_row=n + 1)
        cats = Reference(ws5, min_col=1, min_row=2, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.y_axis.scaling.min, chart.y_axis.scaling.max = 1, 5
        chart.anchor = _two_cell_anchor(m + 3, 1, m + 11, 20)
        ws5.add_chart(chart)

    return wb


@bp.route("/cycles/<int:cycle_id>/subjects/<int:subject_id>/report.xlsx")
@login_required
def export(cycle_id, subject_id):
    if not is_pro():
        abort(403)  # экспорт XLSX — только Pro (раздел 9 ТЗ)
    db = get_db()
    cycle, subj, threshold = _load_target(db, cycle_id, subject_id)
    if cycle["status"] != "closed":
        flash("Отчёт доступен после закрытия цикла.")
        return redirect(url_for("cycles.detail", cycle_id=cycle_id))

    ctx = _assemble(db, cycle, subject_id, subj["employee_id"], threshold, with_dynamics=True)
    wb = _build_workbook(cycle, subj["full_name"], threshold, ctx)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"report_{cycle_id}_{subject_id}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
