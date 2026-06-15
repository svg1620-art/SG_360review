import io
import zipfile

from openpyxl import load_workbook

from tests.helpers import (
    add_assignment,
    add_employee,
    add_subject,
    create_cycle,
    fill_and_submit,
    register,
    set_plan,
)


def _closed_report(app, client):
    register(client)
    set_plan(app, "Acme", "pro")
    subj = add_employee(client, "Субъект")
    ps = [add_employee(client, f"P{i}") for i in range(3)]
    us = [add_employee(client, f"U{i}") for i in range(3)]
    cyc = create_cycle(client, "Q1")
    sb = add_subject(client, cyc, subj)
    for e in ps:
        add_assignment(client, cyc, sb, e, "peer")
    for e in us:
        add_assignment(client, cyc, sb, e, "subordinate")
    client.post(f"/cycles/{cyc}/launch")
    fill_and_submit(app, cyc, lambda rel, ci: 5 if rel == "self" else 2)
    client.post(f"/cycles/{cyc}/close")
    return cyc, sb


def test_export_xlsx_pro(app, client):
    cyc, sb = _closed_report(app, client)
    r = client.get(f"/cycles/{cyc}/subjects/{sb}/report.xlsx")
    assert r.status_code == 200
    assert "spreadsheetml.sheet" in r.headers["Content-Type"]
    data = r.get_data()
    wb = load_workbook(io.BytesIO(data))
    assert {"Сводка", "Я vs другие", "Зоны", "Открытые ответы"}.issubset(set(wb.sheetnames))


def test_export_uses_two_cell_anchor(app, client):
    cyc, sb = _closed_report(app, client)
    data = client.get(f"/cycles/{cyc}/subjects/{sb}/report.xlsx").get_data()
    z = zipfile.ZipFile(io.BytesIO(data))
    two = one = 0
    for name in z.namelist():
        if name.startswith("xl/drawings/") and name.endswith(".xml"):
            xml = z.read(name).decode()
            two += xml.count("twoCellAnchor")
            one += xml.count("oneCellAnchor")
    assert two > 0 and one == 0  # грабли SG: только twoCellAnchor


def test_export_blocked_for_free(app, client):
    cyc, sb = _closed_report(app, client)
    set_plan(app, "Acme", "free")
    assert client.get(f"/cycles/{cyc}/subjects/{sb}/report.xlsx").status_code == 403
